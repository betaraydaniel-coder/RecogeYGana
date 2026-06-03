"""Generate the HOA master database Excel workbook.

The output is intentionally Excel-compatible rather than application-specific:
every major module is stored as an editable table with stable IDs, dropdowns,
starter formulas, and a data dictionary.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName


OUTPUT_FILE = Path(__file__).with_name("HOA_Master_Database_Template.xlsx")


NAVY = "1F4E78"
BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
ORANGE = "F4B183"
LIGHT_ORANGE = "FCE4D6"
GRAY = "808080"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
RED = "C00000"
YELLOW = "FFF2CC"


@dataclass
class WorkbookSheet:
    title: str
    purpose: str
    table_name: str
    columns: List[str]
    validations: Dict[str, str] = field(default_factory=dict)
    formulas: Dict[str, str] = field(default_factory=dict)
    descriptions: Dict[str, str] = field(default_factory=dict)
    required: Iterable[str] = field(default_factory=list)


LOOKUPS: Dict[str, List[str]] = {
    "lkp_state": [
        "AL",
        "AK",
        "AZ",
        "AR",
        "CA",
        "CO",
        "CT",
        "DE",
        "FL",
        "GA",
        "HI",
        "ID",
        "IL",
        "IN",
        "IA",
        "KS",
        "KY",
        "LA",
        "ME",
        "MD",
        "MA",
        "MI",
        "MN",
        "MS",
        "MO",
        "MT",
        "NE",
        "NV",
        "NH",
        "NJ",
        "NM",
        "NY",
        "NC",
        "ND",
        "OH",
        "OK",
        "OR",
        "PA",
        "RI",
        "SC",
        "SD",
        "TN",
        "TX",
        "UT",
        "VT",
        "VA",
        "WA",
        "WV",
        "WI",
        "WY",
        "DC",
        "PR",
        "GU",
        "VI",
        "AS",
        "MP",
    ],
    "lkp_yes_no": ["Yes", "No", "N/A"],
    "lkp_active": ["Active", "Inactive", "Pending", "Archived"],
    "lkp_unit_status": ["Active", "Inactive", "Under Renovation", "Legal Hold"],
    "lkp_property_type": ["Condo", "Townhome", "Single-Family", "Villa", "Common Area", "Other"],
    "lkp_owner_type": ["Individual", "Trust", "LLC", "Corporation", "Government", "Other"],
    "lkp_owner_status": ["Current", "Former", "Pending Transfer"],
    "lkp_contact_type": [
        "Owner",
        "Co-Owner",
        "Tenant",
        "Resident",
        "Board Member",
        "Property Manager",
        "Emergency Contact",
        "Vendor",
        "Attorney",
        "Insurance Agent",
        "Other",
    ],
    "lkp_occupant_type": ["Owner Resident", "Tenant", "Family", "Guest", "Caretaker", "Other"],
    "lkp_billing_frequency": ["Monthly", "Quarterly", "Semiannual", "Annual", "One-Time", "Other"],
    "lkp_statement_delivery": ["Email", "Mail", "Portal", "Email and Mail", "Do Not Send"],
    "lkp_charge_type": [
        "Regular Assessment",
        "Special Assessment",
        "Late Fee",
        "Fine",
        "Interest",
        "Access Card Fee",
        "Move-in Fee",
        "Move-out Fee",
        "Utility Reimbursement",
        "Other",
    ],
    "lkp_charge_status": ["Open", "Paid", "Partial", "Past Due", "Waived", "Void", "In Collections"],
    "lkp_payment_method": [
        "ACH",
        "Check",
        "Credit Card",
        "Debit Card",
        "Cash",
        "Money Order",
        "Wire",
        "Online Portal",
        "Other",
    ],
    "lkp_card_status": ["Available", "Assigned", "Lost", "Stolen", "Disabled", "Returned", "Replaced"],
    "lkp_card_slot": ["1", "2", "Replacement", "Temporary"],
    "lkp_delivery_method": ["In Person", "Mail", "Certified Mail", "Manager Delivery", "Other"],
    "lkp_proof_document": [
        "Deed",
        "Closing Statement",
        "Settlement Statement",
        "Title Commitment",
        "Property Tax Record",
        "Recorded Warranty Deed",
        "Trust/Entity Authority",
        "Other",
    ],
    "lkp_verification_status": ["Pending Review", "Verified", "Rejected", "Expired", "Needs Update"],
    "lkp_document_category": [
        "Ownership Proof",
        "Governing Documents",
        "Financial",
        "Meeting Minutes",
        "Insurance",
        "Vendor Contract",
        "Architectural",
        "Violation",
        "Maintenance",
        "Legal",
        "Reserve",
        "Correspondence",
        "Other",
    ],
    "lkp_confidentiality": [
        "Public to Members",
        "Board Only",
        "Management Only",
        "Confidential",
        "Legal Privileged",
    ],
    "lkp_violation_status": [
        "Open",
        "Notice Sent",
        "Hearing Scheduled",
        "Fined",
        "Cured",
        "Closed",
        "Appealed",
        "Escalated",
    ],
    "lkp_arc_status": [
        "Draft",
        "Submitted",
        "Under Review",
        "Approved",
        "Approved With Conditions",
        "Denied",
        "Withdrawn",
        "Expired",
        "Completed",
    ],
    "lkp_maintenance_status": [
        "New",
        "Assigned",
        "In Progress",
        "Waiting Vendor",
        "Waiting Owner",
        "Completed",
        "Closed",
        "Cancelled",
    ],
    "lkp_priority": ["Low", "Normal", "High", "Urgent"],
    "lkp_meeting_type": ["Board", "Annual", "Special", "Committee", "Executive Session"],
    "lkp_request_status": [
        "Received",
        "In Review",
        "Fulfilled",
        "Denied",
        "Partially Fulfilled",
        "Closed",
    ],
    "lkp_channel": ["Email", "Mail", "Certified Mail", "Phone", "SMS", "Portal", "In Person", "Other"],
    "lkp_policy_status": ["Active", "Expired", "Pending Renewal", "Cancelled"],
    "lkp_asset_status": ["Planned", "Active", "Under Repair", "Retired", "Completed"],
    "lkp_pet_type": ["Dog", "Cat", "Bird", "Fish", "Service Animal", "Other"],
    "lkp_vote_status": ["Eligible", "Ineligible", "Suspended", "Transferred"],
    "lkp_delivery_status": ["Sent", "Delivered", "Bounced", "Returned", "Failed", "Pending"],
}


SHEETS: List[WorkbookSheet] = [
    WorkbookSheet(
        title="Association Profile",
        purpose="Core identity, jurisdiction, management, and policy references for the association.",
        table_name="tblAssociationProfile",
        columns=[
            "Association Name",
            "Legal Entity Name",
            "State",
            "County",
            "Entity Type",
            "Fiscal Year Start",
            "Federal Tax ID Last 4",
            "Registered Agent",
            "Management Company",
            "Records Custodian Contact ID",
            "Operating Bank Last 4",
            "Reserve Bank Last 4",
            "Governing Documents Link",
            "Retention Policy Link",
            "Notes",
        ],
        validations={"State": "lkp_state"},
        required=["Association Name", "State", "Records Custodian Contact ID"],
    ),
    WorkbookSheet(
        title="Units",
        purpose="Master registry for every unit, parcel, common-area asset, or condo interest.",
        table_name="tblUnits",
        columns=[
            "Unit ID",
            "Building",
            "Floor",
            "Unit Number",
            "Legal Description",
            "Parcel/APN",
            "Property Address",
            "Mailing Address",
            "Property Type",
            "Square Feet",
            "Bedrooms",
            "Bathrooms",
            "Parking Space(s)",
            "Storage Unit(s)",
            "Unit Status",
            "Notes",
        ],
        validations={"Property Type": "lkp_property_type", "Unit Status": "lkp_unit_status"},
        required=["Unit ID", "Property Address", "Unit Status"],
    ),
    WorkbookSheet(
        title="Owners",
        purpose="Current and historical ownership records linked to units and ownership proof.",
        table_name="tblOwners",
        columns=[
            "Owner ID",
            "Unit ID",
            "Owner Type",
            "Legal Owner Name",
            "Primary Contact ID",
            "Ownership Start Date",
            "Ownership End Date",
            "Ownership Percentage",
            "Mailing Address",
            "Mailing City",
            "Mailing State",
            "Mailing ZIP",
            "Preferred Billing Delivery",
            "Owner Status",
            "Proof Status",
            "Notes",
        ],
        validations={
            "Owner Type": "lkp_owner_type",
            "Mailing State": "lkp_state",
            "Preferred Billing Delivery": "lkp_statement_delivery",
            "Owner Status": "lkp_owner_status",
        },
        formulas={
            "Proof Status": (
                '=IF([@[Owner ID]]="","",IF(COUNTIFS(tblOwnershipProof[Owner ID],'
                '[@[Owner ID]],tblOwnershipProof[Verification Status],"Verified")>0,'
                '"Verified",IF(COUNTIFS(tblOwnershipProof[Owner ID],[@[Owner ID]])>0,'
                '"Pending/Review","Missing")))'
            )
        },
        required=["Owner ID", "Unit ID", "Legal Owner Name", "Owner Status"],
    ),
    WorkbookSheet(
        title="Directory",
        purpose="Resident, owner, manager, vendor, emergency, and board contact directory.",
        table_name="tblDirectory",
        columns=[
            "Contact ID",
            "Unit ID",
            "Owner ID",
            "Contact Type",
            "Full Name",
            "Relationship to Unit",
            "Primary Email",
            "Secondary Email",
            "Mobile Phone",
            "Home Phone",
            "Work Phone",
            "Preferred Phone",
            "Mailing Address",
            "Emergency Contact Name",
            "Emergency Contact Phone",
            "Communication Consent",
            "Directory Opt-Out",
            "Active",
            "Notes",
        ],
        validations={
            "Contact Type": "lkp_contact_type",
            "Communication Consent": "lkp_yes_no",
            "Directory Opt-Out": "lkp_yes_no",
            "Active": "lkp_active",
        },
        required=["Contact ID", "Full Name", "Contact Type"],
    ),
    WorkbookSheet(
        title="Property Managers",
        purpose="Property manager details for owner-managed or rental units where applicable.",
        table_name="tblPropertyManagers",
        columns=[
            "Manager ID",
            "Unit ID",
            "Company Name",
            "Manager Name",
            "Manager Email",
            "Manager Cell Phone",
            "Office Phone",
            "Authorized Scope",
            "Start Date",
            "End Date",
            "Active",
            "Notes",
        ],
        validations={"Active": "lkp_active"},
        required=["Manager ID", "Unit ID", "Manager Name"],
    ),
    WorkbookSheet(
        title="Occupants & Leases",
        purpose="Occupant, lease, and residency tracking for owners, tenants, guests, and caretakers.",
        table_name="tblOccupantsLeases",
        columns=[
            "Occupant ID",
            "Unit ID",
            "Contact ID",
            "Occupant Type",
            "Move-In Date",
            "Move-Out Date",
            "Lease Start",
            "Lease End",
            "Lease On File",
            "Vehicle Info",
            "Pet Info",
            "Active",
            "Notes",
        ],
        validations={
            "Occupant Type": "lkp_occupant_type",
            "Lease On File": "lkp_yes_no",
            "Active": "lkp_active",
        },
        required=["Occupant ID", "Unit ID", "Contact ID"],
    ),
    WorkbookSheet(
        title="Vehicles",
        purpose="Vehicle, plate, parking permit, and assigned parking tracking.",
        table_name="tblVehicles",
        columns=[
            "Vehicle ID",
            "Unit ID",
            "Contact ID",
            "Plate Number",
            "Plate State",
            "Make",
            "Model",
            "Color",
            "Parking Permit",
            "Assigned Space",
            "Registration Expiration",
            "Active",
            "Notes",
        ],
        validations={"Plate State": "lkp_state", "Active": "lkp_active"},
        required=["Vehicle ID", "Unit ID", "Plate Number"],
    ),
    WorkbookSheet(
        title="Pets",
        purpose="Pet registration and approval records for pet policy administration.",
        table_name="tblPets",
        columns=[
            "Pet ID",
            "Unit ID",
            "Contact ID",
            "Pet Type",
            "Pet Name",
            "Breed",
            "Weight",
            "License/Registration",
            "Vaccination Expiration",
            "Approved",
            "Active",
            "Notes",
        ],
        validations={"Pet Type": "lkp_pet_type", "Approved": "lkp_yes_no", "Active": "lkp_active"},
        required=["Pet ID", "Unit ID", "Pet Type"],
    ),
    WorkbookSheet(
        title="Board & Committees",
        purpose="Board, officer, and committee role records including terms and voting authority.",
        table_name="tblBoardCommittees",
        columns=[
            "Role ID",
            "Contact ID",
            "Role",
            "Committee",
            "Term Start",
            "Term End",
            "Voting Authority",
            "Email",
            "Phone",
            "Active",
            "Notes",
        ],
        validations={"Voting Authority": "lkp_yes_no", "Active": "lkp_active"},
        required=["Role ID", "Contact ID", "Role"],
    ),
    WorkbookSheet(
        title="Access Cards",
        purpose="Assigned access cards, limited to two regular active cards per condo unless replaced or temporary.",
        table_name="tblAccessCards",
        columns=[
            "Card Record ID",
            "Unit ID",
            "Card Slot",
            "Card Number",
            "Assigned Contact ID",
            "Assigned Name",
            "Date Delivered",
            "Delivery Method",
            "Acknowledged By",
            "Signature/Receipt Link",
            "Card Status",
            "Lost/Stolen Date",
            "Disabled Date",
            "Replacement Fee",
            "Notes",
        ],
        validations={
            "Card Slot": "lkp_card_slot",
            "Delivery Method": "lkp_delivery_method",
            "Card Status": "lkp_card_status",
        },
        required=["Card Record ID", "Unit ID", "Card Number", "Card Status"],
    ),
    WorkbookSheet(
        title="Billing Accounts",
        purpose="Billing setup by unit, including assessment amount, frequency, statement delivery, and terms.",
        table_name="tblBillingAccounts",
        columns=[
            "Account ID",
            "Unit ID",
            "Billing Contact ID",
            "Billing Frequency",
            "Regular Assessment Amount",
            "Billing Start Date",
            "Billing End Date",
            "Payment Terms",
            "Auto Late Fee Rule",
            "Statement Delivery",
            "Account Status",
            "Notes",
        ],
        validations={
            "Billing Frequency": "lkp_billing_frequency",
            "Statement Delivery": "lkp_statement_delivery",
            "Account Status": "lkp_active",
        },
        required=["Account ID", "Unit ID", "Billing Frequency"],
    ),
    WorkbookSheet(
        title="Assessment Schedule",
        purpose="Recurring and one-time charge setup approved by the association.",
        table_name="tblAssessmentSchedule",
        columns=[
            "Assessment ID",
            "Applies To",
            "Unit ID/Group",
            "Charge Type",
            "Description",
            "Amount",
            "Frequency",
            "Start Date",
            "End Date",
            "Due Day",
            "Late Fee Amount",
            "Interest Rate",
            "Board Approval Date",
            "Notes",
        ],
        validations={"Charge Type": "lkp_charge_type", "Frequency": "lkp_billing_frequency"},
        required=["Assessment ID", "Charge Type", "Amount", "Frequency"],
    ),
    WorkbookSheet(
        title="Charges",
        purpose="All billed amounts: assessments, special assessments, fines, fees, interest, and adjustments.",
        table_name="tblCharges",
        columns=[
            "Charge ID",
            "Account ID",
            "Unit ID",
            "Owner ID",
            "Charge Type",
            "Description",
            "Charge Date",
            "Due Date",
            "Amount",
            "Payment Applied",
            "Open Balance",
            "Charge Status",
            "Related Record ID",
            "Memo",
            "Created By",
        ],
        validations={"Charge Type": "lkp_charge_type", "Charge Status": "lkp_charge_status"},
        formulas={"Open Balance": '=IF([@[Charge ID]]="","",MAX(0,[@Amount]-[@[Payment Applied]]))'},
        required=["Charge ID", "Account ID", "Unit ID", "Amount", "Due Date"],
    ),
    WorkbookSheet(
        title="Payments",
        purpose="Payment receipt, deposit, and application details for owner accounts.",
        table_name="tblPayments",
        columns=[
            "Payment ID",
            "Account ID",
            "Unit ID",
            "Owner ID",
            "Payment Date",
            "Payment Amount",
            "Payment Method",
            "Reference Number",
            "Deposit Account",
            "Applied Charge ID",
            "Applied Amount",
            "Unapplied Amount",
            "Received By",
            "Receipt Sent",
            "Notes",
        ],
        validations={"Payment Method": "lkp_payment_method", "Receipt Sent": "lkp_yes_no"},
        formulas={
            "Unapplied Amount": '=IF([@[Payment ID]]="","",MAX(0,[@[Payment Amount]]-[@[Applied Amount]]))'
        },
        required=["Payment ID", "Account ID", "Unit ID", "Payment Amount", "Payment Date"],
    ),
    WorkbookSheet(
        title="Billing Ledger",
        purpose="Audit-style transaction ledger for charges, payments, credits, and adjustments.",
        table_name="tblBillingLedger",
        columns=[
            "Ledger ID",
            "Account ID",
            "Unit ID",
            "Transaction Date",
            "Transaction Type",
            "Source Record ID",
            "Debit",
            "Credit",
            "Running Balance",
            "Posted By",
            "Notes",
        ],
        formulas={
            "Running Balance": (
                '=IF([@[Ledger ID]]="","",SUMIFS(tblBillingLedger[Debit],tblBillingLedger[Account ID],'
                '[@[Account ID]],tblBillingLedger[Transaction Date],"<="&[@[Transaction Date]])-'
                'SUMIFS(tblBillingLedger[Credit],tblBillingLedger[Account ID],[@[Account ID]],'
                'tblBillingLedger[Transaction Date],"<="&[@[Transaction Date]]))'
            )
        },
        required=["Ledger ID", "Account ID", "Unit ID", "Transaction Date"],
    ),
    WorkbookSheet(
        title="Delinquency Tracking",
        purpose="Aging, notices, collections, payment plans, attorney, and lien workflow tracking.",
        table_name="tblDelinquencyTracking",
        columns=[
            "Delinquency ID",
            "Unit ID",
            "Account ID",
            "As Of Date",
            "Balance Due",
            "Oldest Due Date",
            "Days Past Due",
            "Collection Stage",
            "Last Notice Date",
            "Next Action Date",
            "Attorney/Lien Status",
            "Payment Plan",
            "Notes",
        ],
        formulas={
            "Balance Due": '=IF([@[Unit ID]]="","",SUMIFS(tblCharges[Open Balance],tblCharges[Unit ID],[@[Unit ID]]))',
            "Days Past Due": '=IF([@[Oldest Due Date]]="","",MAX(0,[@[As Of Date]]-[@[Oldest Due Date]]))',
        },
        required=["Delinquency ID", "Unit ID", "As Of Date"],
    ),
    WorkbookSheet(
        title="Ownership Proof",
        purpose="Valid proof of ownership registration and verification record for each owner/unit.",
        table_name="tblOwnershipProof",
        columns=[
            "Proof ID",
            "Unit ID",
            "Owner ID",
            "Document Type",
            "Document Date",
            "Recording Number",
            "Issuing County/Agency",
            "Document Link/Location",
            "Received Date",
            "Verification Status",
            "Verified By",
            "Verification Date",
            "Expiration/Review Date",
            "Confidentiality",
            "Notes",
        ],
        validations={
            "Document Type": "lkp_proof_document",
            "Verification Status": "lkp_verification_status",
            "Confidentiality": "lkp_confidentiality",
        },
        required=["Proof ID", "Unit ID", "Owner ID", "Document Type", "Verification Status"],
    ),
    WorkbookSheet(
        title="Documents",
        purpose="Central index for governing documents, financials, minutes, contracts, insurance, and correspondence.",
        table_name="tblDocuments",
        columns=[
            "Document ID",
            "Category",
            "Related Unit ID",
            "Related Owner/Contact ID",
            "Title",
            "Document Date",
            "Version",
            "Storage Location/Link",
            "Retention Period",
            "Confidentiality",
            "Uploaded By",
            "Upload Date",
            "Review/Expiration Date",
            "Notes",
        ],
        validations={"Category": "lkp_document_category", "Confidentiality": "lkp_confidentiality"},
        required=["Document ID", "Category", "Title", "Storage Location/Link"],
    ),
    WorkbookSheet(
        title="Violations",
        purpose="Compliance and rule enforcement records with notice, cure, hearing, fine, and appeal fields.",
        table_name="tblViolations",
        columns=[
            "Violation ID",
            "Unit ID",
            "Reported Date",
            "Rule/CCR Reference",
            "Category",
            "Description",
            "Photo/Document Link",
            "Notice Date",
            "Cure Deadline",
            "Hearing Date",
            "Fine Amount",
            "Charge ID",
            "Status",
            "Closed Date",
            "Notes",
        ],
        validations={"Status": "lkp_violation_status"},
        required=["Violation ID", "Unit ID", "Reported Date", "Description", "Status"],
    ),
    WorkbookSheet(
        title="Architectural Requests",
        purpose="ARC/ACC request workflow for exterior changes, approvals, conditions, and completion verification.",
        table_name="tblArchitecturalRequests",
        columns=[
            "ARC ID",
            "Unit ID",
            "Owner ID",
            "Submitted Date",
            "Project Type",
            "Description",
            "Contractor",
            "Documents Link",
            "Review Deadline",
            "Board/Committee Decision",
            "Decision Date",
            "Conditions",
            "Completion Due",
            "Completion Verified Date",
            "Status",
            "Notes",
        ],
        validations={"Status": "lkp_arc_status"},
        required=["ARC ID", "Unit ID", "Submitted Date", "Description", "Status"],
    ),
    WorkbookSheet(
        title="Maintenance",
        purpose="Work orders, common-area maintenance, vendor assignments, estimates, approvals, and completion.",
        table_name="tblMaintenance",
        columns=[
            "Request ID",
            "Unit ID/Common Area",
            "Reported By Contact ID",
            "Reported Date",
            "Location",
            "Issue Type",
            "Description",
            "Priority",
            "Assigned Vendor ID",
            "Estimate Amount",
            "Approved Amount",
            "Work Order Link",
            "Status",
            "Completed Date",
            "Notes",
        ],
        validations={"Priority": "lkp_priority", "Status": "lkp_maintenance_status"},
        required=["Request ID", "Reported Date", "Description", "Priority", "Status"],
    ),
    WorkbookSheet(
        title="Vendors",
        purpose="Vendor, contractor, insurance, W-9, preferred status, and contract lifecycle tracking.",
        table_name="tblVendors",
        columns=[
            "Vendor ID",
            "Company Name",
            "Contact Name",
            "Email",
            "Phone",
            "Service Category",
            "Contract Start",
            "Contract End",
            "Insurance Expiration",
            "W-9 On File",
            "Preferred",
            "Active",
            "Notes",
        ],
        validations={"W-9 On File": "lkp_yes_no", "Preferred": "lkp_yes_no", "Active": "lkp_active"},
        required=["Vendor ID", "Company Name", "Service Category"],
    ),
    WorkbookSheet(
        title="Meetings",
        purpose="Board, annual, committee, and executive-session meeting records and retained minutes.",
        table_name="tblMeetings",
        columns=[
            "Meeting ID",
            "Meeting Type",
            "Date",
            "Location/Platform",
            "Agenda Link",
            "Minutes Link",
            "Approval Date",
            "Quorum Met",
            "Motions/Resolutions",
            "Record Retention",
            "Confidentiality",
            "Notes",
        ],
        validations={
            "Meeting Type": "lkp_meeting_type",
            "Quorum Met": "lkp_yes_no",
            "Confidentiality": "lkp_confidentiality",
        },
        required=["Meeting ID", "Meeting Type", "Date"],
    ),
    WorkbookSheet(
        title="Communications",
        purpose="Announcements, notices, mailings, opt-out tracking, and related-record communications.",
        table_name="tblCommunications",
        columns=[
            "Communication ID",
            "Date",
            "Audience",
            "Channel",
            "Subject",
            "Related Unit ID",
            "Related Record ID",
            "Sent By",
            "Delivery Status",
            "Opt-Out Respected",
            "Attachment Link",
            "Notes",
        ],
        validations={
            "Channel": "lkp_channel",
            "Delivery Status": "lkp_delivery_status",
            "Opt-Out Respected": "lkp_yes_no",
        },
        required=["Communication ID", "Date", "Audience", "Subject"],
    ),
    WorkbookSheet(
        title="Insurance",
        purpose="Association insurance policies, carrier, broker, coverage, deductibles, and renewal tracking.",
        table_name="tblInsurance",
        columns=[
            "Policy ID",
            "Carrier",
            "Broker Contact ID",
            "Policy Type",
            "Policy Number",
            "Coverage Limit",
            "Deductible",
            "Effective Date",
            "Expiration Date",
            "Certificate Link",
            "Renewal Owner",
            "Status",
            "Notes",
        ],
        validations={"Status": "lkp_policy_status"},
        required=["Policy ID", "Carrier", "Policy Type", "Expiration Date"],
    ),
    WorkbookSheet(
        title="Reserves & Assets",
        purpose="Reserve components, capital projects, inspection dates, funding priority, and replacement costs.",
        table_name="tblReservesAssets",
        columns=[
            "Asset/Project ID",
            "Asset Type",
            "Location",
            "Description",
            "Useful Life",
            "Last Inspection Date",
            "Next Inspection Date",
            "Reserve Component",
            "Estimated Replacement Cost",
            "Current Reserve Balance",
            "Funding Priority",
            "Status",
            "Notes",
        ],
        validations={"Funding Priority": "lkp_priority", "Status": "lkp_asset_status"},
        required=["Asset/Project ID", "Asset Type", "Description"],
    ),
    WorkbookSheet(
        title="Elections & Voting",
        purpose="Voting interests, proxies, eligibility, ballots, and election administration.",
        table_name="tblElectionsVoting",
        columns=[
            "Voting Record ID",
            "Election ID",
            "Unit ID",
            "Owner ID",
            "Voting Interest",
            "Voting Status",
            "Proxy Holder Contact ID",
            "Proxy Received Date",
            "Ballot Received Date",
            "Ballot Valid",
            "Notes",
        ],
        validations={"Voting Status": "lkp_vote_status", "Ballot Valid": "lkp_yes_no"},
        required=["Voting Record ID", "Election ID", "Unit ID"],
    ),
    WorkbookSheet(
        title="Record Requests",
        purpose="Owner/member inspection and copy requests, deadlines, fees, fulfillment, and confidentiality review.",
        table_name="tblRecordRequests",
        columns=[
            "Request ID",
            "Requester Contact ID",
            "Unit ID",
            "Request Date",
            "Request Method",
            "Records Requested",
            "Statutory/Policy Deadline",
            "Fulfillment Date",
            "Fees Charged",
            "Fee Paid",
            "Status",
            "Confidentiality Review",
            "Notes",
        ],
        validations={
            "Request Method": "lkp_channel",
            "Fee Paid": "lkp_yes_no",
            "Status": "lkp_request_status",
            "Confidentiality Review": "lkp_confidentiality",
        },
        required=["Request ID", "Requester Contact ID", "Request Date", "Records Requested", "Status"],
    ),
    WorkbookSheet(
        title="Unit Compliance",
        purpose="Formula-based control sheet for card limits, verified ownership proof, balances, and open workflows.",
        table_name="tblUnitCompliance",
        columns=[
            "Unit ID",
            "Active Card Count",
            "Card Limit Check",
            "Current Owner Count",
            "Verified Proof Count",
            "Open Balance",
            "Open Violations",
            "Pending ARC Requests",
            "Notes",
        ],
        formulas={
            "Active Card Count": '=IF([@[Unit ID]]="","",COUNTIFS(tblAccessCards[Unit ID],[@[Unit ID]],tblAccessCards[Card Status],"Assigned"))',
            "Card Limit Check": '=IF([@[Unit ID]]="","",IF([@[Active Card Count]]<=2,"OK","Over 2 Assigned Cards"))',
            "Current Owner Count": '=IF([@[Unit ID]]="","",COUNTIFS(tblOwners[Unit ID],[@[Unit ID]],tblOwners[Owner Status],"Current"))',
            "Verified Proof Count": '=IF([@[Unit ID]]="","",COUNTIFS(tblOwnershipProof[Unit ID],[@[Unit ID]],tblOwnershipProof[Verification Status],"Verified"))',
            "Open Balance": '=IF([@[Unit ID]]="","",SUMIFS(tblCharges[Open Balance],tblCharges[Unit ID],[@[Unit ID]]))',
            "Open Violations": '=IF([@[Unit ID]]="","",COUNTIFS(tblViolations[Unit ID],[@[Unit ID]],tblViolations[Status],"<>Closed",tblViolations[Status],"<>Cured"))',
            "Pending ARC Requests": '=IF([@[Unit ID]]="","",COUNTIFS(tblArchitecturalRequests[Unit ID],[@[Unit ID]],tblArchitecturalRequests[Status],"Submitted")+COUNTIFS(tblArchitecturalRequests[Unit ID],[@[Unit ID]],tblArchitecturalRequests[Status],"Under Review"))',
        },
        required=["Unit ID"],
    ),
]


FIELD_DESCRIPTIONS = {
    "Unit ID": "Stable identifier for the condo/unit. Use the same value across related sheets.",
    "Owner ID": "Stable identifier for an owner or ownership entity.",
    "Contact ID": "Stable identifier for a person or company in the directory.",
    "Account ID": "Billing account identifier, usually one per unit unless local policy requires otherwise.",
    "Document Link/Location": "File path, cloud link, document management reference, or storage box location.",
    "Storage Location/Link": "File path, cloud link, document management reference, or storage box location.",
    "Confidentiality": "Visibility classification; confirm with governing documents and state law before disclosure.",
    "Verification Status": "Review status for proof of ownership documentation.",
    "Open Balance": "Formula field based on charges and applied payments.",
    "Card Limit Check": "Flags units with more than two assigned active access cards.",
    "Statutory/Policy Deadline": "Deadline from state law, governing documents, or board policy.",
}


def create_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "HOA Master Database Template"
    wb.properties.subject = "Excel-compatible all-in-one HOA database"
    wb.properties.creator = "Cursor Cloud Agent"
    wb.properties.keywords = "HOA, condominium, Excel, database, billing, access cards, ownership proof"

    create_instructions_sheet(wb)
    create_dashboard_sheet(wb)
    create_research_coverage_sheet(wb)
    create_lookup_sheet(wb)
    for sheet in SHEETS:
        create_table_sheet(wb, sheet)
    create_data_dictionary_sheet(wb)
    order_sheets(wb)
    return wb


def create_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HOA Master Database Template"
    ws["A1"].font = Font(bold=True, size=22, color=NAVY)
    ws["A2"] = "Excel-compatible, editable, all-in-one workbook for HOA and condominium association operations."
    ws["A2"].font = Font(italic=True, color=GRAY)

    sections = [
        (
            "Purpose",
            "This workbook centralizes owner/resident directory data, property manager details, assigned access cards, billing, billing records, proof of ownership, governance records, maintenance, compliance, documents, reserves, and record requests.",
        ),
        (
            "How to use",
            "Start with Association Profile and Units, then add Owners, Directory contacts, Ownership Proof, Billing Accounts, Access Cards, and any operational records. Use consistent IDs across sheets.",
        ),
        (
            "Excel tables",
            "Each worksheet uses a structured Excel table. Add new rows directly below the first entry row; formulas and dropdowns will carry forward in modern Excel.",
        ),
        (
            "Access card control",
            "Access Cards records each assigned card. Unit Compliance flags any unit with more than two Assigned cards, while preserving replacement and temporary card history.",
        ),
        (
            "Billing model",
            "Billing Accounts defines account setup, Assessment Schedule defines recurring or approved charges, Charges records amounts billed, Payments records money received, and Billing Ledger supports audit-style bookkeeping.",
        ),
        (
            "Ownership proof",
            "Ownership Proof tracks document type, recording information, storage location, verification status, reviewer, and confidentiality classification.",
        ),
        (
            "U.S. HOA coverage",
            "The included modules reflect common U.S. HOA/condo management systems and record categories: accounting, assessments, owner rosters, minutes, governing documents, contracts, violations, architectural review, maintenance, insurance, reserves, elections, and records requests. State-specific legal requirements vary; have counsel or the association manager confirm local retention, inspection, and privacy rules.",
        ),
        (
            "Data protection",
            "Limit access to this workbook. It can contain personal contact information, financial records, ownership documents, legal notes, and other confidential material.",
        ),
    ]

    row = 4
    for heading, body in sections:
        ws.cell(row=row, column=1, value=heading)
        ws.cell(row=row, column=1).font = Font(bold=True, color=NAVY)
        ws.cell(row=row + 1, column=1, value=body)
        ws.cell(row=row + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
        row += 3

    ws["A29"] = "Recommended ID prefixes"
    ws["A29"].font = Font(bold=True, color=NAVY)
    prefixes = [
        ("Unit", "U-001 or B1-101"),
        ("Owner", "OWN-001"),
        ("Contact", "CON-001"),
        ("Access Card", "CARD-001"),
        ("Billing Account", "ACCT-001"),
        ("Charge", "CHG-001"),
        ("Payment", "PAY-001"),
        ("Document", "DOC-001"),
    ]
    for offset, (name, example) in enumerate(prefixes, start=30):
        ws.cell(row=offset, column=1, value=name)
        ws.cell(row=offset, column=2, value=example)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    for col in range(3, 7):
        ws.column_dimensions[get_column_letter(col)].width = 18


def create_dashboard_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HOA Operations Dashboard"
    ws["A1"].font = Font(bold=True, size=20, color=NAVY)
    ws["A2"] = "Live summary formulas populate as data is entered into workbook tables."
    ws["A2"].font = Font(italic=True, color=GRAY)

    headers = ["Metric", "Value", "Notes"]
    metrics = [
        ("Total Units", '=COUNTIF(tblUnits[Unit ID],"<>")', "Registered units/common interests"),
        ("Active Units", '=COUNTIFS(tblUnits[Unit Status],"Active")', "Units currently active"),
        ("Current Owners", '=COUNTIFS(tblOwners[Owner Status],"Current")', "Current ownership records"),
        ("Directory Contacts", '=COUNTIF(tblDirectory[Contact ID],"<>")', "All contact records"),
        ("Assigned Access Cards", '=COUNTIFS(tblAccessCards[Card Status],"Assigned")', "Currently assigned cards"),
        ("Units Over Card Limit", '=COUNTIFS(tblUnitCompliance[Card Limit Check],"Over 2 Assigned Cards")', "Review immediately"),
        ("Open Charges", '=COUNTIFS(tblCharges[Open Balance],">0")', "Charges with unpaid balance"),
        ("Outstanding Balance", '=SUM(tblCharges[Open Balance])', "Total open balance"),
        (
            "Payments This Year",
            '=SUMIFS(tblPayments[Payment Amount],tblPayments[Payment Date],">="&DATE(YEAR(TODAY()),1,1),tblPayments[Payment Date],"<="&TODAY())',
            "Cash receipts dated in current year",
        ),
        (
            "Past Due Charges",
            '=COUNTIFS(tblCharges[Due Date],"<"&TODAY(),tblCharges[Open Balance],">0")',
            "Open balances past due date",
        ),
        (
            "Pending Ownership Proof",
            '=COUNTIFS(tblOwnershipProof[Verification Status],"Pending Review")+COUNTIFS(tblOwnershipProof[Verification Status],"Needs Update")',
            "Documents awaiting review",
        ),
        ("Open Violations", '=COUNTIFS(tblViolations[Status],"<>Closed",tblViolations[Status],"<>Cured",tblViolations[Violation ID],"<>")', "Open enforcement records"),
        (
            "ARC Requests Under Review",
            '=COUNTIFS(tblArchitecturalRequests[Status],"Submitted")+COUNTIFS(tblArchitecturalRequests[Status],"Under Review")',
            "Architectural requests needing action",
        ),
        ("Open Maintenance", '=COUNTIFS(tblMaintenance[Status],"<>Closed",tblMaintenance[Status],"<>Completed",tblMaintenance[Request ID],"<>")', "Maintenance/work orders not closed"),
        ("Upcoming Insurance Renewals", '=COUNTIFS(tblInsurance[Expiration Date],"<="&TODAY()+60,tblInsurance[Status],"Active")', "Active policies expiring within 60 days"),
        ("Open Record Requests", '=COUNTIFS(tblRecordRequests[Status],"<>Closed",tblRecordRequests[Status],"<>Fulfilled",tblRecordRequests[Request ID],"<>")', "Inspection/copy requests requiring action"),
    ]

    start_row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(metrics, start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row_idx, column=2).number_format = '#,##0.00;[Red]-#,##0.00;0'

    table_ref = f"A{start_row}:C{start_row + len(metrics)}"
    table = Table(displayName="tblDashboard", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 58

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Key Operational Counts"
    chart.y_axis.title = "Metric"
    chart.x_axis.title = "Value"
    data = Reference(ws, min_col=2, min_row=5, max_row=10)
    cats = Reference(ws, min_col=1, min_row=5, max_row=10)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 15
    ws.add_chart(chart, "E4")

    ws.conditional_formatting.add(
        "B10:B20",
        FormulaRule(formula=["B10>0"], fill=PatternFill("solid", fgColor=YELLOW)),
    )


def create_research_coverage_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Research Coverage")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "U.S. HOA Database Coverage Matrix"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = (
        "This matrix maps common U.S. HOA and condominium management database needs "
        "to the worksheets and controls included in this template."
    )
    ws["A2"].font = Font(italic=True, color=GRAY)

    headers = ["Area", "Included Worksheets", "Key Data Captured", "Why It Matters"]
    rows = [
        (
            "Owner, resident, and property directory",
            "Units; Owners; Directory; Property Managers; Occupants & Leases",
            "Units, legal owners, residents, emails, mobile phones, emergency contacts, manager email and cell phone.",
            "Supports notices, billing, emergency response, member lists, manager coordination, and unit-level accountability.",
        ),
        (
            "Access control",
            "Access Cards; Unit Compliance",
            "Two regular card slots per condo, card numbers, delivery date, acknowledgement link, lost/stolen/disabled status.",
            "Creates an auditable record of physical access credentials and flags units exceeding the two-card rule.",
        ),
        (
            "Billing and accounting",
            "Billing Accounts; Assessment Schedule; Charges; Payments; Billing Ledger; Delinquency Tracking",
            "Assessments, fees, fines, payments, references, balances, ledger activity, notices, collection stage, lien/attorney notes.",
            "Reflects common HOA software expectations for dues collection, owner ledgers, aging, and financial transparency.",
        ),
        (
            "Proof of ownership",
            "Ownership Proof; Documents; Owners",
            "Deeds and other valid proof, recording details, document location, verification status, reviewer, review date.",
            "Helps confirm owner standing, voting rights, billing responsibility, records access eligibility, and transfer history.",
        ),
        (
            "Governance and official records",
            "Board & Committees; Meetings; Elections & Voting; Documents; Record Requests",
            "Roles, terms, minutes, agendas, motions, ballots, proxies, retained documents, owner inspection requests, deadlines.",
            "Covers common association records such as minutes, governing documents, elections, policies, and owner record requests.",
        ),
        (
            "Compliance and enforcement",
            "Violations; Charges; Communications; Documents",
            "CCR/rule references, notices, photos, cure dates, hearings, fines, appeals, closure, related communication.",
            "Provides a consistent enforcement trail and links financial penalties to billing records.",
        ),
        (
            "Architectural review",
            "Architectural Requests; Documents; Communications",
            "Submissions, project details, contractor, review deadline, decision, conditions, completion verification.",
            "Mirrors common ARC/ACC workflows used by U.S. HOA and condominium associations.",
        ),
        (
            "Maintenance and vendor management",
            "Maintenance; Vendors; Documents; Insurance",
            "Work orders, priorities, estimates, approvals, vendors, insurance expirations, contracts, W-9 status.",
            "Supports common-area operations, vendor accountability, and contract/insurance review.",
        ),
        (
            "Risk, insurance, reserves, and assets",
            "Insurance; Reserves & Assets",
            "Policies, limits, deductibles, renewal dates, reserve components, inspections, useful life, replacement cost.",
            "Helps boards monitor risk exposure and capital reserve planning.",
        ),
        (
            "Communications and privacy",
            "Communications; Directory; Documents; Record Requests",
            "Channels, delivery status, opt-out handling, confidentiality classification, related records, attachments.",
            "Supports consistent notice history while recognizing state-specific privacy and disclosure limits.",
        ),
        (
            "Vehicles, parking, and pets",
            "Vehicles; Pets; Occupants & Leases",
            "Plates, permits, spaces, registration expiration, pet type, approvals, vaccine/license dates.",
            "Useful for condo rules, parking enforcement, amenity access, and resident policy administration.",
        ),
    ]

    start_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    table_ref = f"A{start_row}:D{start_row + len(rows)}"
    table = Table(displayName="tblResearchCoverage", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 62


def create_lookup_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lookups")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Lookup values for dropdowns"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Edit these lists carefully; data validation ranges are named ranges used across the workbook."
    ws["A2"].font = Font(italic=True, color=GRAY)

    for col_idx, (name, values) in enumerate(LOOKUPS.items(), start=1):
        col_letter = get_column_letter(col_idx)
        ws.cell(row=4, column=col_idx, value=name)
        ws.cell(row=4, column=col_idx).font = Font(bold=True, color=WHITE)
        ws.cell(row=4, column=col_idx).fill = PatternFill("solid", fgColor=NAVY)
        for row_offset, value in enumerate(values, start=5):
            ws.cell(row=row_offset, column=col_idx, value=value)
        ws.column_dimensions[col_letter].width = max(18, min(34, max(len(name), *(len(v) for v in values)) + 2))
        ref = f"{quote_sheetname(ws.title)}!${col_letter}$5:${col_letter}${4 + len(values)}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    ws.freeze_panes = "A5"


def create_table_sheet(wb: Workbook, sheet: WorkbookSheet) -> None:
    ws = wb.create_sheet(sheet.title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = sheet.title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = sheet.purpose
    ws["A2"].font = Font(italic=True, color=GRAY)

    header_row = 4
    first_data_row = 5
    for col_idx, column in enumerate(sheet.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=column)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if column in sheet.required:
            add_required_comment(ws, header_row, col_idx)

        data_cell = ws.cell(row=first_data_row, column=col_idx)
        if column in sheet.formulas:
            data_cell.value = sheet.formulas[column]
        data_cell.alignment = Alignment(wrap_text=True, vertical="top")

    ref = f"A{header_row}:{get_column_letter(len(sheet.columns))}{first_data_row}"
    table = Table(displayName=sheet.table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True)
    ws.add_table(table)

    for column, lookup_name in sheet.validations.items():
        if column not in sheet.columns:
            continue
        col_idx = sheet.columns.index(column) + 1
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(type="list", formula1=f"={lookup_name}", allow_blank=True)
        dv.error = "Select a value from the dropdown list."
        dv.errorTitle = "Invalid value"
        dv.prompt = "Choose from the approved lookup values."
        dv.promptTitle = column
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}{first_data_row}:{col_letter}1000")

    for col_idx, column in enumerate(sheet.columns, start=1):
        width = min(max(len(column) + 4, 14), 34)
        if any(token in column.lower() for token in ["description", "notes", "address", "link", "requested"]):
            width = 36
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = ref
    apply_sheet_formats(ws, sheet)


def add_required_comment(ws, row: int, col: int) -> None:
    cell = ws.cell(row=row, column=col)
    cell.fill = PatternFill("solid", fgColor=GREEN)


def apply_sheet_formats(ws, sheet: WorkbookSheet) -> None:
    currency_headers = {
        "Regular Assessment Amount",
        "Amount",
        "Late Fee Amount",
        "Fine Amount",
        "Replacement Fee",
        "Payment Amount",
        "Applied Amount",
        "Unapplied Amount",
        "Debit",
        "Credit",
        "Running Balance",
        "Balance Due",
        "Fees Charged",
        "Estimate Amount",
        "Approved Amount",
        "Coverage Limit",
        "Deductible",
        "Estimated Replacement Cost",
        "Current Reserve Balance",
        "Open Balance",
    }
    percent_headers = {"Ownership Percentage", "Interest Rate", "Voting Interest"}
    date_keywords = ["Date", "Start", "End", "Expiration", "Deadline", "Due", "Received", "Effective", "Review"]

    for col_idx, header in enumerate(sheet.columns, start=1):
        col_letter = get_column_letter(col_idx)
        data_range = f"{col_letter}5:{col_letter}1000"
        for row in range(5, 21):
            cell = ws.cell(row=row, column=col_idx)
            if header in currency_headers:
                cell.number_format = '$#,##0.00;[Red]-$#,##0.00;'
            elif header in percent_headers:
                cell.number_format = '0.00%'
            elif any(keyword in header for keyword in date_keywords):
                cell.number_format = 'm/d/yyyy'

        if header in currency_headers:
            ws.conditional_formatting.add(
                data_range,
                FormulaRule(formula=[f'{col_letter}5<0'], fill=PatternFill("solid", fgColor=LIGHT_ORANGE)),
            )

    if sheet.title == "Unit Compliance":
        card_col = sheet.columns.index("Card Limit Check") + 1
        col_letter = get_column_letter(card_col)
        ws.conditional_formatting.add(
            f"{col_letter}5:{col_letter}1000",
            FormulaRule(
                formula=[f'{col_letter}5="Over 2 Assigned Cards"'],
                fill=PatternFill("solid", fgColor=LIGHT_ORANGE),
                font=Font(color=RED, bold=True),
            ),
        )


def create_data_dictionary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Dictionary")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Data Dictionary"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Field-level reference for workbook users and future database migration."
    ws["A2"].font = Font(italic=True, color=GRAY)

    headers = ["Worksheet", "Table", "Field", "Required", "Dropdown/Formula", "Description"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")

    row = 5
    for sheet in SHEETS:
        for column in sheet.columns:
            dropdown_or_formula = ""
            if column in sheet.validations:
                dropdown_or_formula = sheet.validations[column]
            if column in sheet.formulas:
                dropdown_or_formula = "Formula"
            description = (
                sheet.descriptions.get(column)
                or FIELD_DESCRIPTIONS.get(column)
                or infer_description(sheet.title, column)
            )
            values = [
                sheet.title,
                sheet.table_name,
                column,
                "Yes" if column in sheet.required else "No",
                dropdown_or_formula,
                description,
            ]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row, column=col_idx, value=value)
                ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

    ref = f"A4:F{row - 1}"
    table = Table(displayName="tblDataDictionary", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 72


def infer_description(sheet_name: str, column: str) -> str:
    clean = column.replace("/", " or ")
    if clean.endswith(" ID"):
        return f"Identifier used to link this {sheet_name.lower()} record to related workbook tables."
    if "Date" in clean or "Start" in clean or "End" in clean or "Deadline" in clean:
        return f"Relevant date for {clean.lower()}."
    if "Status" in clean or clean == "Active":
        return f"Current workflow or lifecycle status for this {sheet_name.lower()} record."
    if "Email" in clean:
        return "Email address for notices, billing, or association communication."
    if "Phone" in clean:
        return "Phone number for operational or emergency contact."
    if "Notes" in clean:
        return "Additional context, exceptions, board notes, or management notes."
    return f"{clean} for the {sheet_name.lower()} record."


def order_sheets(wb: Workbook) -> None:
    desired = ["Instructions", "Dashboard", "Research Coverage", "Lookups"] + [sheet.title for sheet in SHEETS] + ["Data Dictionary"]
    wb._sheets.sort(key=lambda ws: desired.index(ws.title) if ws.title in desired else len(desired))


def finalize_workbook(wb: Workbook) -> None:
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                alignment = copy(cell.alignment)
                alignment.wrap_text = cell.alignment.wrap_text
                alignment.vertical = cell.alignment.vertical or "top"
                cell.alignment = alignment
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.freeze_panes = ws.freeze_panes
    wb.active = 0


def main() -> None:
    wb = create_workbook()
    finalize_workbook(wb)
    wb.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
