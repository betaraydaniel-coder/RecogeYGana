"""Generate a maximum-compatibility HOA workbook.

This variant intentionally avoids Excel tables, charts, conditional formatting,
and structured-reference formulas. It is meant for users opening the workbook
in older Excel versions, Apple Numbers, Google Sheets, or mobile viewers.
"""

from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from generate_hoa_workbook import FIELD_DESCRIPTIONS, LOOKUPS, SHEETS, infer_description


OUTPUT_FILE = Path(__file__).with_name("HOA_Master_Database_Template_Compatible.xlsx")

NAVY = "1F4E78"
WHITE = "FFFFFF"
GRAY = "808080"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def create_compatible_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    wb.properties.title = "HOA Master Database Template - Compatible"
    wb.properties.subject = "Excel-compatible all-in-one HOA database"
    wb.properties.creator = "Cursor Cloud Agent"

    create_instructions(ws)
    create_dashboard(wb)
    create_lookup_sheet(wb)
    create_research_coverage(wb)
    for sheet in SHEETS:
        create_plain_sheet(wb, sheet)
    create_data_dictionary(wb)
    finalize_workbook(wb)
    return wb


def create_instructions(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HOA Master Database Template - Compatible Edition"
    ws["A1"].font = Font(bold=True, size=20, color=NAVY)
    ws["A2"] = "Plain Excel workbook for reliable editing across Excel, Google Sheets, Numbers, and mobile viewers."
    ws["A2"].font = Font(italic=True, color=GRAY)

    sections = [
        (
            "What this file includes",
            "A complete HOA and condominium association database structure covering directory, owners, property managers, access cards, billing, billing records, ownership proof, documents, violations, architectural requests, maintenance, vendors, meetings, insurance, reserves, elections, vehicles, pets, communications, and record requests.",
        ),
        (
            "Compatibility approach",
            "This edition uses plain worksheets with headers, frozen panes, dropdowns, and formatting. It avoids advanced Excel table objects, charts, and structured-reference formulas so it opens more reliably in different spreadsheet programs.",
        ),
        (
            "How to start",
            "Enter units first, then owners, directory contacts, ownership proof, billing accounts, and access cards. Use the same IDs across related sheets.",
        ),
        (
            "Access card rule",
            "Use Access Cards to record two standard cards per condo, card number, assigned person, delivery date, delivery method, acknowledgement link, and status.",
        ),
        (
            "Legal and privacy note",
            "State HOA and condominium requirements vary. Confirm retention, inspection, privacy, collection, and disclosure requirements with qualified counsel or the association manager.",
        ),
    ]

    row = 4
    for heading, body in sections:
        ws.cell(row=row, column=1, value=heading).font = Font(bold=True, color=NAVY)
        ws.cell(row=row + 1, column=1, value=body).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
        row += 3

    ws["A22"] = "Recommended ID prefixes"
    ws["A22"].font = Font(bold=True, color=NAVY)
    prefixes = [
        ("Unit", "U-001 or B1-101"),
        ("Owner", "OWN-001"),
        ("Contact", "CON-001"),
        ("Access Card", "CARD-001"),
        ("Billing Account", "ACCT-001"),
        ("Charge", "CHG-001"),
        ("Payment", "PAY-001"),
        ("Document", "DOC-001"),
    ]
    for offset, (name, example) in enumerate(prefixes, start=23):
        ws.cell(row=offset, column=1, value=name)
        ws.cell(row=offset, column=2, value=example)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30
    for col in range(3, 7):
        ws.column_dimensions[get_column_letter(col)].width = 18


def create_dashboard(wb: Workbook) -> None:
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HOA Operations Dashboard"
    ws["A1"].font = Font(bold=True, size=18, color=NAVY)
    ws["A2"] = "Enter values manually or add formulas that fit your spreadsheet program."
    ws["A2"].font = Font(italic=True, color=GRAY)

    headers = ["Metric", "Value", "Notes"]
    rows = [
        ("Total Units", "", "Count records in Units."),
        ("Current Owners", "", "Count current owner records."),
        ("Directory Contacts", "", "Count contacts in Directory."),
        ("Assigned Access Cards", "", "Count cards with status Assigned."),
        ("Units Over Card Limit", "", "Review units with more than two active assigned cards."),
        ("Open Charges", "", "Count unpaid charges."),
        ("Outstanding Balance", "", "Sum unpaid charge balances."),
        ("Past Due Charges", "", "Count open balances past due."),
        ("Pending Ownership Proof", "", "Count proof records pending review or update."),
        ("Open Violations", "", "Count unresolved enforcement records."),
        ("ARC Requests Under Review", "", "Count submitted or under-review requests."),
        ("Open Maintenance", "", "Count not-completed maintenance records."),
        ("Upcoming Insurance Renewals", "", "Policies expiring soon."),
        ("Open Record Requests", "", "Owner/member record requests requiring action."),
    ]
    write_header(ws, headers, row=4)
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 58


def create_lookup_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lookups")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Lookup values for dropdowns"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Edit carefully; dropdowns reference these named ranges."
    ws["A2"].font = Font(italic=True, color=GRAY)

    for col_idx, (name, values) in enumerate(LOOKUPS.items(), start=1):
        col_letter = get_column_letter(col_idx)
        ws.cell(row=4, column=col_idx, value=name)
        ws.cell(row=4, column=col_idx).font = Font(bold=True, color=WHITE)
        ws.cell(row=4, column=col_idx).fill = PatternFill("solid", fgColor=NAVY)
        for row_offset, value in enumerate(values, start=5):
            ws.cell(row=row_offset, column=col_idx, value=value)
        ws.column_dimensions[col_letter].width = max(18, min(34, max(len(name), *(len(v) for v in values)) + 2))
        ref = f"{quote_sheetname(ws.title)}!${col_letter}$5:${col_letter}${4 + len(values)}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))
    ws.freeze_panes = "A5"


def create_research_coverage(wb: Workbook) -> None:
    ws = wb.create_sheet("Research Coverage")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "U.S. HOA Database Coverage Matrix"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    headers = ["Area", "Included Worksheets", "Key Data Captured"]
    rows = [
        ("Directory", "Units; Owners; Directory; Property Managers", "Emails, phones, mobile phones, manager contact details, emergency contacts."),
        ("Access Cards", "Access Cards", "Two card slots per condo, card number, delivery date, method, acknowledgement, status."),
        ("Billing", "Billing Accounts; Assessment Schedule; Charges; Payments; Billing Ledger; Delinquency Tracking", "Assessments, charges, payments, balances, references, aging, collections."),
        ("Proof of Ownership", "Ownership Proof; Documents; Owners", "Deed/proof type, recording details, document location, verification status."),
        ("Governance", "Board & Committees; Meetings; Elections & Voting; Record Requests", "Roles, minutes, agendas, voting, records inspection requests."),
        ("Compliance", "Violations; Architectural Requests; Communications", "Notices, CCR references, ARC decisions, deadlines, closure."),
        ("Operations", "Maintenance; Vendors; Insurance; Reserves & Assets", "Work orders, contracts, insurance renewals, reserve components, assets."),
    ]
    write_header(ws, headers, row=4)
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 76


def create_plain_sheet(wb: Workbook, sheet) -> None:
    ws = wb.create_sheet(sheet.title)
    ws.sheet_view.showGridLines = False
    ws["A1"] = sheet.title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = sheet.purpose
    ws["A2"].font = Font(italic=True, color=GRAY)
    write_header(ws, sheet.columns, row=4, required=set(sheet.required))
    for col_idx, column in enumerate(sheet.columns, start=1):
        width = min(max(len(column) + 4, 14), 34)
        if any(token in column.lower() for token in ["description", "notes", "address", "link", "requested"]):
            width = 38
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        if column in sheet.validations:
            col_letter = get_column_letter(col_idx)
            dv = DataValidation(type="list", formula1=f"={sheet.validations[column]}", allow_blank=True)
            dv.error = "Select a value from the dropdown list."
            dv.errorTitle = "Invalid value"
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}5:{col_letter}1000")
    for row in range(5, 25):
        for col_idx in range(1, len(sheet.columns) + 1):
            ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(sheet.columns))}1000"


def create_data_dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Dictionary")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Data Dictionary"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    headers = ["Worksheet", "Field", "Required", "Dropdown", "Description"]
    write_header(ws, headers, row=4)
    row = 5
    for sheet in SHEETS:
        for column in sheet.columns:
            description = sheet.descriptions.get(column) or FIELD_DESCRIPTIONS.get(column) or infer_description(sheet.title, column)
            values = [
                sheet.title,
                column,
                "Yes" if column in sheet.required else "No",
                sheet.validations.get(column, ""),
                description,
            ]
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row, column=col_idx, value=value)
                ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 74


def write_header(ws, headers, row: int, required=None) -> None:
    required = required or set()
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor="70AD47" if header in required else NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def finalize_workbook(wb: Workbook) -> None:
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                alignment = copy(cell.alignment)
                alignment.vertical = cell.alignment.vertical or "top"
                cell.alignment = alignment
        for row_idx in range(5, min(ws.max_row + 1, 1001)):
            if row_idx % 2 == 0:
                for col_idx in range(1, ws.max_column + 1):
                    if ws.cell(row=row_idx, column=col_idx).value is None:
                        ws.cell(row=row_idx, column=col_idx).fill = fill
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.active = 0


def main() -> None:
    wb = create_compatible_workbook()
    wb.save(OUTPUT_FILE)
    print(f"Created {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
